"""Compara dos Excel celda por celda (valor y formato) FUERA de las zonas de la app.

Es la prueba de ida y vuelta de la Fase 3 vista con otra librería: la app escribe con ExcelJS y
aquí se lee con openpyxl, la misma que generó el Excel real. Compara:
- nombres, orden y visibilidad de las hojas que ya existían;
- en cada hoja que ya existía, todas las celdas fuera de las columnas de la app: valor, fuente,
  relleno, alineación, bordes, formato de número y protección;
- anchos de columna (fuera de la zona), altos de fila, paneles inmovilizados, celdas combinadas,
  validaciones de datos y formato condicional.

Uso:
  .venv/bin/python scripts/comparar_excel.py ORIGINAL NUEVO --zonas '{"GR2QB": [10, 11, 12]}' \
      --hojas-app "Asistencia (app)" "Detalle (app)" "_app"
Imprime un JSON {"diferencias": [...], "celdas_comparadas": n}. Sale con 1 si hay diferencias.
No imprime el contenido de las celdas, solo su dirección y qué cambió (el Excel real tiene datos personales).
"""

import argparse
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter


def color(c):
    if c is None:
        return None
    if c.type == "rgb":
        # El canal alfa no se usa en rellenos ni fuentes: openpyxl escribe 00, otras librerías FF.
        return ("rgb", str(c.rgb)[-6:])
    if c.type == "theme":
        return ("theme", c.theme, round(c.tint or 0, 6))
    if c.type == "indexed":
        return ("indexed", c.indexed)
    return (c.type,)


def estilo(c):
    f, fi, a, b, p = c.font, c.fill, c.alignment, c.border, c.protection
    relleno = (fi.fill_type, color(fi.fgColor) if fi.fill_type else None) if fi is not None else None
    return {
        "fuente": (f.name, f.sz, bool(f.b), bool(f.i), f.u, bool(f.strike), color(f.color)),
        "relleno": relleno,
        "alineacion": (a.horizontal, a.vertical, bool(a.wrap_text), a.indent or 0, a.text_rotation or 0),
        "bordes": tuple((s.style, color(s.color) if s.style else None) for s in (b.left, b.right, b.top, b.bottom)),
        "formato": c.number_format,
        "proteccion": (p.locked, p.hidden),
    }


def anchos(ws):
    """Ancho de cada columna (n.° → ancho); una definición puede cubrir varias columnas (min–max)."""
    salida = {}
    for dim in ws.column_dimensions.values():
        if dim.customWidth and dim.width is not None and dim.min:
            for col in range(dim.min, (dim.max or dim.min) + 1):
                salida[col] = dim.width
    return salida


def valor(c):
    v = c.value
    if isinstance(v, str) and v == "":
        return None
    return v


def comparar(ruta_a, ruta_b, zonas, hojas_app, maximo=200):
    a = openpyxl.load_workbook(ruta_a)
    b = openpyxl.load_workbook(ruta_b)
    difs = []
    comparadas = 0

    def dif(hoja, donde, que):
        if len(difs) < maximo:
            difs.append({"hoja": hoja, "donde": donde, "que": que})

    propias = set(hojas_app)
    hojas_a = [ws.title for ws in a.worksheets if ws.title not in propias]
    hojas_b = [ws.title for ws in b.worksheets if ws.title not in propias]
    if hojas_a != hojas_b:
        dif("*", "libro", f"hojas distintas: {hojas_a} → {hojas_b}")

    for nombre in hojas_a:
        if nombre not in b.sheetnames:
            continue
        wa, wb = a[nombre], b[nombre]
        zona = set(zonas.get(nombre, []))
        if wa.sheet_state != wb.sheet_state:
            dif(nombre, "hoja", "visibilidad")
        if wa.freeze_panes != wb.freeze_panes:
            dif(nombre, "hoja", "paneles inmovilizados")
        if sorted(map(str, wa.merged_cells.ranges)) != sorted(map(str, wb.merged_cells.ranges)):
            dif(nombre, "hoja", "celdas combinadas")
        if len(wa.data_validations.dataValidation) != len(wb.data_validations.dataValidation):
            dif(nombre, "hoja", "validaciones de datos")
        if len(list(wa.conditional_formatting)) != len(list(wb.conditional_formatting)):
            dif(nombre, "hoja", "formato condicional")

        max_fila = max(wa.max_row, wb.max_row)
        max_col = max(wa.max_column, wb.max_column)
        anchos_a, anchos_b = anchos(wa), anchos(wb)
        for col in sorted(set(anchos_a) | set(anchos_b)):
            if col not in zona and anchos_a.get(col) != anchos_b.get(col):
                dif(nombre, f"columna {get_column_letter(col)}", "ancho")
        for fila in range(1, max_fila + 1):
            ha = wa.row_dimensions[fila].height if fila in wa.row_dimensions else None
            hb = wb.row_dimensions[fila].height if fila in wb.row_dimensions else None
            if ha != hb:
                dif(nombre, f"fila {fila}", "alto")
            for col in range(1, max_col + 1):
                if col in zona:
                    continue
                ca, cb = wa.cell(row=fila, column=col), wb.cell(row=fila, column=col)
                comparadas += 1
                if valor(ca) != valor(cb):
                    dif(nombre, ca.coordinate, "valor")
                if ca.has_style or cb.has_style:
                    ea, eb = estilo(ca), estilo(cb)
                    for k in ea:
                        if ea[k] != eb[k]:
                            dif(nombre, ca.coordinate, f"formato: {k}")
    return {"diferencias": difs, "celdas_comparadas": comparadas}


def main():
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("original")
    p.add_argument("nuevo")
    p.add_argument("--zonas", default="{}", help='JSON {"HOJA": [columnas de la app (números)]}')
    p.add_argument("--hojas-app", nargs="*", default=[])
    args = p.parse_args()
    r = comparar(args.original, args.nuevo, json.loads(args.zonas), args.hojas_app)
    print(json.dumps(r, ensure_ascii=False, indent=1))
    sys.exit(1 if r["diferencias"] else 0)


if __name__ == "__main__":
    main()
