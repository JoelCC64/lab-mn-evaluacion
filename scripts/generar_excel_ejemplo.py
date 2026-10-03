"""Genera los Excel ficticios de /datos-ejemplo con la forma exacta de los reales (Anexo J).

- Cursos_Lab_MN_2026B_EJEMPLO.xlsx: misma estructura y formato que Cursos_Lab_MN_2026B.xlsx
  (hoja Resumen + una hoja por paralelo; encabezados en la fila 4; fila ámbar para pendientes).
- Asistencia_Semana1_EJEMPLO.xlsx: misma estructura que Asistencia_Semana1_Lab_MN_2026B.xlsx.

Todos los nombres, códigos y correos son INVENTADOS (códigos 1999xxxxx, correos @epn.example).
Se usa openpyxl, la misma librería que generó el Excel real, para que las pruebas de lectura y de
escritura (Fase 3) trabajen sobre un archivo con la misma estructura interna.

Uso:  .venv/bin/python scripts/generar_excel_ejemplo.py
"""

import json
import random
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "datos-ejemplo"
AZUL, BLANCO, BANDA, ROJO, VERDE, AMBAR = "1F3A5F", "FFFFFF", "EEF3F8", "FDE2E2", "E3F4E1", "FFF4D6"

# Alumnos por curso (de la nómina; no es un dato personal) y pendientes ficticios.
MATRICULADOS = {"GR2QB": 22, "GR9EB": 19, "GR3MB": 20, "GR1AA": 20, "GR4EB": 18,
                "GR6CD": 20, "GR7SA": 18, "GR7EB": 19, "GR3QA": 21, "GR1QA": 17}
PENDIENTES = {"GR4EB": 1, "GR3QA": 1}
SIN_GRUPO = {"GR9EB": 3, "GR3MB": 2, "GR1AA": 1, "GR7SA": 2, "GR1QA": 3, "GR3QA": 1}
SIN_CORREO = {"GR2QB"}                      # como el «directo a carrera» real
SIN_CLASE_SEMANA1 = {"GR2QB": "Sin clase (permiso de las autoridades)"}
DIAS = {"lunes": "Lunes", "martes": "Martes", "miercoles": "Miércoles", "jueves": "Jueves", "viernes": "Viernes"}
METODOLOGIA = {"TRAD": "Clásica", "SQI": "SQI"}

APELLIDOS = """ACOSTA AGUILAR ALBÁN ALMEIDA ÁLVAREZ ANDRADE ARIAS ARMIJOS BALSECA BARRAGÁN BENÍTEZ BRAVO
CABRERA CALDERÓN CARRERA CASTILLO CASTRO CEVALLOS CHÁVEZ CHICAIZA CISNEROS CÓRDOVA CORREA CRUZ DÁVILA
DELGADO DÍAZ ENRÍQUEZ ESPÍN ESPINOSA ESTRELLA FIALLOS FLORES GALARZA GALLEGOS GARCÉS GARZÓN GUAMÁN
GUERRERO GUEVARA HERRERA HIDALGO IZURIETA JARAMILLO JIMÉNEZ LARA LEÓN LOZADA MALDONADO MEDINA MEJÍA
MENA MERINO MOLINA MONTALVO MORALES MOREIRA MOSQUERA NARANJO NAVARRETE NÚÑEZ OCHOA ORDÓÑEZ ORTEGA
PACHECO PAREDES PAZMIÑO PEÑAFIEL POVEDA PROAÑO QUINTANA RAMÍREZ RAMOS REYES RIVADENEIRA ROBALINO
ROMERO RUIZ SALAZAR SALGADO SALINAS SANDOVAL SANTANA SARMIENTO SILVA SOLÍS SUÁREZ TAPIA TERÁN TOAPANTA
TORRES TRUJILLO URBINA VACA VALENCIA VALLEJO VARGAS VÁSCONEZ VELASCO VILLACÍS VILLAGÓMEZ YÁNEZ ZAMBRANO
ZAPATA ZURITA""".split()
NOMBRES_F = """ALEJANDRA ALISSON ANA ANDREA CAMILA CAROLINA CRISTINA DANIELA DIANA DOMÉNICA ELIZABETH
EMILIA EVELYN FERNANDA GABRIELA ISABEL JOHANNA KAREN LUCÍA MARÍA MICAELA NATALIA NICOLE PAOLA SAMANTHA
SOFÍA TAMARA VALENTINA VALERIA VANESSA XIMENA""".split()
NOMBRES_M = """ADRIÁN ALEJANDRO ANDRÉS ÁNGEL ARIEL BRYAN CARLOS DANIEL DAVID DIEGO EDUARDO ESTEBAN
FERNANDO GABRIEL ISAAC IVÁN JAVIER JHON JORGE JOSÉ JOSUÉ JUAN KEVIN LEONARDO LUIS MARCO MARTÍN MATEO
MIGUEL NICOLÁS PABLO PAÚL RAFAEL RICARDO SANTIAGO SEBASTIÁN STEVEN""".split()


def sin_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def norm(s):
    return sin_tildes(s).lower()


def cursos():
    datos = json.loads((RAIZ / "config" / "cursos-2026B.json").read_text(encoding="utf-8"))
    return datos["cursos"]


def estudiantes_ficticios(rng):
    """Filas ficticias por curso, con grupos de 3 y 4, algunos sin grupo y los pendientes al final."""
    usados, codigo = set(), 199900100
    por_curso = {}
    for c in cursos():
        p = c["paralelo"]
        n_oficiales, n_pend = MATRICULADOS[p], PENDIENTES.get(p, 0)
        filas = []
        for k in range(n_oficiales + n_pend):
            while True:
                a1, a2 = rng.sample(APELLIDOS, 2)
                n1, n2 = rng.sample(rng.choice([NOMBRES_F, NOMBRES_M]), 2)
                nombre = f"{a1} {a2} {n1} {n2}"
                if nombre not in usados:
                    usados.add(nombre)
                    break
            codigo += rng.randint(3, 97)
            correo = None if p in SIN_CORREO else f"{norm(n1)}.{norm(a1)}{rng.randint(10, 99)}@epn.example"
            filas.append({"codigo": codigo, "nombre": nombre, "correo": correo,
                          "pendiente": k >= n_oficiales, "grupo": None, "observacion": None})
        oficiales = sorted((f for f in filas if not f["pendiente"]), key=lambda f: norm(f["nombre"]))
        pendientes = [f for f in filas if f["pendiente"]]
        for f in pendientes:
            f["observacion"] = "Pendiente de la lista final: no consta en la nómina (registrado con el QR; dato ficticio)."
        # Grupos de 3 y 4 entre quienes tienen grupo (en orden aleatorio, como el registro por QR).
        con_grupo = rng.sample(oficiales + pendientes, len(oficiales) + len(pendientes) - SIN_GRUPO.get(p, 0))
        g, i = 1, 0
        while i < len(con_grupo):
            resto = len(con_grupo) - i
            tam = 4 if resto % 3 != 0 and resto >= 4 else 3
            for f in con_grupo[i:i + tam]:
                f["grupo"] = g
            g, i = g + 1, i + tam
        if p == "GR4EB":
            oficiales[2]["observacion"] = "Solicitó cambio de horario en la primera semana (dato ficticio)."
        por_curso[p] = oficiales + pendientes
    return por_curso


def titulo(ws, texto, subtitulo):
    ws["A1"] = texto
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = subtitulo
    ws["A2"].font = Font(italic=True, color="5A6270")


def encabezado(ws, fila, columnas, anchos=None):
    for j, col in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=j, value=col)
        c.font = Font(bold=True, color=BLANCO)
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if anchos:
            ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    if anchos:
        ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def pintar(ws, fila0, filas, ncol, color_extra=None):
    banda = False
    for i, f in enumerate(filas):
        banda = not banda
        fondo = AMBAR if f["pendiente"] else (BLANCO if banda else BANDA)
        for j in range(1, ncol + 1):
            ws.cell(row=fila0 + i, column=j).fill = PatternFill("solid", fgColor=fondo)
        if color_extra:
            col, color = color_extra(f)
            if color:
                ws.cell(row=fila0 + i, column=col).fill = PatternFill("solid", fgColor=color)


def resumen(ws, por_curso, titulo_texto, subtitulo):
    titulo(ws, titulo_texto, subtitulo)
    cols = ["Paralelo", "Día", "Horario", "Grupo A/B", "Metodología", "Laboratorio", "Aula (talleres)",
            "Estudiantes", "Pendientes", "Fuente"]
    encabezado(ws, 4, cols, [11, 12, 14, 10, 12, 12, 15, 12, 13, 20])
    fila = 5
    for c in cursos():
        p = c["paralelo"]
        vals = [p, DIAS[c["dia"]], f"{c['inicio']}–{c['fin']}", c["cronograma"], METODOLOGIA[c["metodologia"]],
                c.get("laboratorio"), c.get("aula"), MATRICULADOS[p], PENDIENTES.get(p) or None, "Ficticio"]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=fila, column=j, value=v)
        fila += 1
    ws.cell(row=fila, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=fila, column=8, value=sum(MATRICULADOS.values())).font = Font(bold=True)
    ws.cell(row=fila, column=9, value=sum(PENDIENTES.values())).font = Font(bold=True)


def excel_cursos(por_curso):
    wb = openpyxl.Workbook()
    resumen(wb.active, por_curso, "Laboratorio de Mecánica Newtoniana 2026B · Cursos propios (EJEMPLO)",
            "DATOS FICTICIOS para pruebas: ningún nombre, código ni correo corresponde a un estudiante real.")
    wb.active.title = "Resumen"
    for c in cursos():
        p = c["paralelo"]
        filas = por_curso[p]
        n_pend = PENDIENTES.get(p, 0)
        ws = wb.create_sheet(p)
        extra = (f" + {n_pend} pendiente{'s' if n_pend > 1 else ''} de la lista final: "
                 "no consta en la nómina (fila ámbar)") if n_pend else ""
        titulo(ws, f"{p} · {DIAS[c['dia']]} {c['inicio']}–{c['fin']} · Grupo {c['cronograma']} · "
                   f"{METODOLOGIA[c['metodologia']]} · {c.get('laboratorio')} · aula {c.get('aula')}",
               f"{MATRICULADOS[p]} estudiantes{extra}")
        cols = ["N°", "Código único", "Apellidos y nombres", "Correo institucional", "Grupo A/B", "Metodología",
                "Grupo de trabajo", "Observación"]
        encabezado(ws, 4, cols, [5, 13, 42, 36, 10, 12, 10, 50])
        for i, f in enumerate(filas, start=5):
            vals = [i - 4, f["codigo"], f["nombre"], f["correo"], c["cronograma"], METODOLOGIA[c["metodologia"]],
                    f["grupo"], f["observacion"]]
            for j, v in enumerate(vals, start=1):
                ws.cell(row=i, column=j, value=v)
        pintar(ws, 5, filas, len(cols))
    destino = SALIDA / "Cursos_Lab_MN_2026B_EJEMPLO.xlsx"
    wb.save(destino)
    return destino


def excel_asistencia(por_curso):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    titulo(ws, "Asistencia · Semana 1 (28 sep – 1 oct 2026) · Laboratorio de Mecánica Newtoniana (EJEMPLO)",
           "DATOS FICTICIOS. Asistió = registrado en un grupo de trabajo mediante el QR de su paralelo.")
    encabezado(ws, 4, ["Paralelo", "Día", "Horario", "Grupo A/B", "Metodología", "Matriculados", "Asistieron",
                       "No asistieron", "% asistencia", "Estado de la clase"], [11, 12, 14, 10, 12, 13, 11, 13, 13, 38])
    wc = wb.create_sheet("Correos")
    titulo(wc, "Correos para enviar el material", "Copiar la celda y pegar en «Para» o «CCO» de Outlook.")
    encabezado(wc, 4, ["Paralelo", "Grupo A/B", "Metodología", "Estado de la clase", "Todos", "N°"], [10, 10, 12, 26, 60, 5])
    fila_r = 5
    for k, c in enumerate(cursos()):
        p = c["paralelo"]
        filas = por_curso[p]
        sin_clase = SIN_CLASE_SEMANA1.get(p)

        def estado(f):
            if sin_clase:
                return sin_clase
            if f["grupo"] is None:
                return "No asistió"
            return "Asistió · pendiente (no consta en la nómina)" if f["pendiente"] else "Asistió"

        oficiales = [f for f in filas if not f["pendiente"]]
        asist = sum(1 for f in oficiales if estado(f) == "Asistió")
        vals = [p, DIAS[c["dia"]], f"{c['inicio']}–{c['fin']}", c["cronograma"], METODOLOGIA[c["metodologia"]],
                len(oficiales), None if sin_clase else asist, None if sin_clase else len(oficiales) - asist,
                "—" if sin_clase else f"{round(100 * asist / len(oficiales))} %", sin_clase or "Realizada"]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=fila_r, column=j, value=v)
        fila_r += 1
        correos = "; ".join(f["correo"] for f in filas if f["correo"])
        for j, v in enumerate([p, c["cronograma"], METODOLOGIA[c["metodologia"]], sin_clase or "Realizada",
                               correos or None, sum(1 for f in filas if f["correo"])], start=1):
            wc.cell(row=5 + k, column=j, value=v)

        wp = wb.create_sheet(p)
        titulo(wp, f"{p} · {DIAS[c['dia']]} {c['inicio']}–{c['fin']} · Grupo {c['cronograma']} · {METODOLOGIA[c['metodologia']]}",
               f"Estado de la clase: {sin_clase or 'Realizada'}")
        cols = ["N°", "Código único", "Apellidos y nombres", "Grupo A/B", "Metodología", "Asistencia",
                "Grupo de trabajo", "Correo institucional", "Observación"]
        encabezado(wp, 4, cols, [5, 13, 40, 10, 12, 32, 10, 36, 50])
        for i, f in enumerate(filas, start=5):
            v = [i - 4, f["codigo"], f["nombre"], c["cronograma"], METODOLOGIA[c["metodologia"]], estado(f),
                 f["grupo"], f["correo"], f["observacion"]]
            for j, x in enumerate(v, start=1):
                wp.cell(row=i, column=j, value=x)
        color = lambda f: (6, VERDE if estado(f) == "Asistió" else ROJO if estado(f) == "No asistió" else AMBAR)
        pintar(wp, 5, filas, len(cols), color_extra=color)
    wo = wb.create_sheet("Observaciones")
    titulo(wo, "Observaciones de la semana 1", "Sin observaciones (datos ficticios).")
    destino = SALIDA / "Asistencia_Semana1_EJEMPLO.xlsx"
    wb.save(destino)
    return destino


def main():
    SALIDA.mkdir(exist_ok=True)
    por_curso = estudiantes_ficticios(random.Random(2026))
    for d in (excel_cursos(por_curso), excel_asistencia(por_curso)):
        print("Generado:", d.relative_to(RAIZ))


if __name__ == "__main__":
    main()
